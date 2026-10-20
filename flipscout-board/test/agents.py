#!/usr/bin/env python3
"""The listing agents' brokerages, for the Test board's agent cards.

    python3 -I agents.py <prod sheet .xlsx> <out agents.js>

data.js keeps only the agent's name: the sheet's "Listing Agent" column is
"Name, Brokerage" ("Michael Freethy, Compass") and build_data.py cuts it at
the first comma. This writes what comes after that comma, keyed by MLS #:

    window.__AGENTS__ = {"v":1, "built":"<ISO>", "bk":{"<MLS #>":"Compass", ...}};

The board groups leads into agents itself (by email, then phone, then name);
this file only adds the brokerage. A lead missing from it shows none.
"""
import datetime, json, re, sys
import openpyxl

MLS = re.compile(r'^[A-Z0-9]{5,20}$')


def main():
    src, out = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    if 'Leads' not in wb.sheetnames:
        sys.exit('the sheet has no Leads tab')
    rows = wb['Leads'].iter_rows(values_only=True)
    head = [str(h or '').strip() for h in next(rows)]
    if 'MLS #' not in head or 'Listing Agent' not in head:
        sys.exit('the Leads tab has no "MLS #" or "Listing Agent" column')
    im, ia = head.index('MLS #'), head.index('Listing Agent')
    bk = {}
    for r in rows:
        r = list(r) + [None] * (len(head) - len(r))
        mls = str(r[im] or '').strip().upper()
        la = re.sub(r'\s+', ' ', str(r[ia] or '')).strip()
        if not MLS.match(mls) or ',' not in la:
            continue
        b = la.split(',', 1)[1].strip(' ,')
        # a licence number, or a stray "$3", is not a brokerage
        if b and re.search(r'[A-Za-z]', b) and not re.match(r'^(?:dre|lic|cal ?bre|#)\b', b, re.I):
            bk[mls] = b[:80]
    body = {'v': 1, 'built': datetime.datetime.now(datetime.timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ'), 'bk': bk}
    with open(out, 'w', encoding='utf8') as f:
        f.write('window.__AGENTS__ = ' + json.dumps(body, ensure_ascii=False, separators=(',', ':')) + ';\n')
    print('agents.js: brokerage for %d leads' % len(bk))


if __name__ == '__main__':
    main()
