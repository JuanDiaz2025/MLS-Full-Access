#!/usr/bin/env python3
"""The ARV and the sold homes it rests on, for the Test board's comps panel.

    python3 -I comps.py <prod sheet .xlsx> <out comps.js>

The FlipScout app (v1.73+) writes five columns at the end of the Leads tab:
"ARV", "ARV $/SqFt", "Comp Confidence", "Comps" (compact JSON, one entry per
sold home) and "Comps Run On". This writes them out keyed by MLS #:

    window.__COMPS__ = {"v":1, "built":"<ISO>", "by":{"<MLS #>":
        {"arv":1200000, "ppsf":800, "conf":"High", "at":"2026-10-20",
         "list":[{"mls","addr","price","sqft","ppsf","sold","bd","ba","yr","use","why"}, ...]}}};

A lead without an ARV is left out (its Comp Confidence says why on the
sheet); the board then shows "no comps yet". Rehab, holding, selling costs
and profit are the board's own math, from its Settings — not in this file.
Every field is checked, as core.boardComps does in the app: a cell typed by
hand, or cut short, never reaches the page as anything but a number, a date
or plain text.
"""
import datetime, json, re, sys
import openpyxl

MLS = re.compile(r'^[A-Z0-9]{5,20}$')
DAY = re.compile(r'^\d{4}-\d{2}-\d{2}$')
LIST_MAX = 16
CONF = ('High', 'Medium', 'Low')


def num(v):
    """A sheet cell as a number: 1200000, "1,200,000", "$800" → float, else None."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r'[^0-9.]', '', str(v))
    try:
        return float(s) if s else None
    except ValueError:
        return None


def tidy(v):
    """A count that may be fractional (2.5 baths): 3.0 → 3, 2.5 stays."""
    n = num(v)
    return None if n is None else int(n) if n == int(n) else n


def whole(v):
    n = num(v)
    return int(round(n)) if n is not None else None


def text(v, n):
    return re.sub(r'[\x00-\x1f\x7f  ]+', ' ', str(v or '')).strip()[:n]


def day(v):
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.strftime('%Y-%m-%d')
    s = str(v or '').strip()[:10]
    m = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{2,4})$', s)
    if m:
        y = m.group(3) if len(m.group(3)) == 4 else '20' + m.group(3)
        s = '%s-%02d-%02d' % (y, int(m.group(1)), int(m.group(2)))
    return s if DAY.match(s) else ''


def comp_list(cell):
    try:
        a = json.loads(cell) if isinstance(cell, str) and cell.strip() else []
    except ValueError:
        return []
    out = []
    for c in (a if isinstance(a, list) else [])[:LIST_MAX]:
        if not isinstance(c, dict):
            continue
        mls = text(c.get('mls'), 20).upper()
        if not MLS.match(mls):
            continue
        out.append({
            'mls': mls, 'addr': text(c.get('addr'), 120),
            'price': whole(c.get('price')), 'sqft': whole(c.get('sqft')), 'ppsf': whole(c.get('ppsf')),
            'sold': day(c.get('sold')), 'bd': tidy(c.get('bd')), 'ba': tidy(c.get('ba')), 'yr': whole(c.get('yr')),
            'use': 1 if c.get('use') else 0, 'why': text(c.get('why'), 80),
        })
    return out


def main():
    src, out = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    if 'Leads' not in wb.sheetnames:
        sys.exit('the sheet has no Leads tab')
    rows = wb['Leads'].iter_rows(values_only=True)
    head = [str(h or '').strip() for h in next(rows)]
    if 'MLS #' not in head or 'ARV' not in head:
        sys.exit('the Leads tab has no "MLS #" or "ARV" column (comps need FlipScout v1.73+)')
    col = lambda name: head.index(name) if name in head else None
    im, ia, ip, ic, il, it = (col(n) for n in ('MLS #', 'ARV', 'ARV $/SqFt', 'Comp Confidence', 'Comps', 'Comps Run On'))
    get = lambda r, i: r[i] if i is not None else None
    by, skipped = {}, 0
    for r in rows:
        r = list(r) + [None] * (len(head) - len(r))
        mls = str(r[im] or '').strip().upper()
        if not MLS.match(mls):
            continue
        arv = whole(get(r, ia))
        if not arv or arv < 50000:
            skipped += 1
            continue
        conf = text(get(r, ic), 20)
        by[mls] = {
            'arv': arv, 'ppsf': whole(get(r, ip)),
            'conf': conf if conf in CONF else '',
            'at': day(get(r, it)),
            'list': comp_list(get(r, il)),
        }
    body = {'v': 1, 'built': datetime.datetime.now(datetime.timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ'), 'by': by}
    with open(out, 'w', encoding='utf8') as f:
        f.write('window.__COMPS__ = ' + json.dumps(body, ensure_ascii=False, separators=(',', ':')) + ';\n')
    print('comps.js: ARV for %d leads (%d without one)' % (len(by), skipped))


if __name__ == '__main__':
    main()
